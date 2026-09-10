from __future__ import annotations

import io
import html
from datetime import datetime
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak


def excel_report_bytes(clean_df, quality, profile, kpis, insights, recs) -> bytes:
    out=io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        clean_df.to_excel(writer,index=False,sheet_name="Cleaned Data")
        q=pd.DataFrame([{k:v for k,v in quality.items() if not isinstance(v,(dict,list))}]); q.to_excel(writer,index=False,sheet_name="Quality Summary")
        pd.DataFrame([{"column":c,**m} for c,m in profile.items()]).to_excel(writer,index=False,sheet_name="Data Dictionary")
        pd.DataFrame(kpis).drop(columns=[c for c in ["raw"] if c in pd.DataFrame(kpis).columns]).to_excel(writer,index=False,sheet_name="KPIs")
        pd.DataFrame(insights).to_excel(writer,index=False,sheet_name="Insights")
        pd.DataFrame(recs).to_excel(writer,index=False,sheet_name="Recommendations")
        for ws in writer.book.worksheets:
            ws.freeze_panes="A2"
            ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]:
                cell.font=Font(bold=True,color="FFFFFF")
                cell.fill=PatternFill("solid",fgColor="272A56")
                cell.alignment=Alignment(horizontal="center")
            for col_cells in ws.columns:
                max_len=min(50,max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)+2)
                ws.column_dimensions[get_column_letter(col_cells[0].column)].width=max_len
    return out.getvalue()


def pdf_report_bytes(title, clean_df, quality, kpis, insights, recs) -> bytes:
    out=io.BytesIO(); doc=SimpleDocTemplate(out,pagesize=A4,rightMargin=36,leftMargin=36,topMargin=40,bottomMargin=40)
    styles=getSampleStyleSheet(); styles.add(ParagraphStyle(name="Small",parent=styles["BodyText"],fontSize=9,leading=12,textColor=colors.HexColor("#475569")))
    story=[Paragraph(html.escape(str(title)),styles["Title"]),Paragraph(f"Generated {datetime.now():%d %b %Y, %H:%M}",styles["Small"]),Spacer(1,12)]
    story += [Paragraph("Data quality",styles["Heading2"])]
    qdata=[["Metric","Value"],["Rows",f"{quality['rows']:,}"],["Columns",f"{quality['columns']:,}"],["Missing cells",f"{quality['missing']:,}"],["Duplicates",f"{quality['duplicates']:,}"],["Invalid values",f"{quality['invalid']:,}"],["Outliers",f"{quality['outliers']:,}"],["Quality score",f"{quality['score']}/100"]]
    t=Table(qdata,colWidths=[220,100]); t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#272A56")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),.25,colors.HexColor("#CBD5E1")),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F8FAFC")])]))
    story += [t,Spacer(1,12),Paragraph("Key metrics",styles["Heading2"])]
    kdata=[["KPI","Value","Description"]]+[[k.get("label",""),k.get("value",""),k.get("sub","")] for k in kpis]
    kt=Table(kdata,colWidths=[180,110,140]); kt.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#272A56")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),.25,colors.HexColor("#CBD5E1")),("VALIGN",(0,0),(-1,-1),"TOP")]))
    story += [kt,Spacer(1,12),Paragraph("Insights",styles["Heading2"])]
    for i in insights: story.append(Paragraph(f"<b>{html.escape(str(i.get('title','Insight')))}</b>: {html.escape(str(i.get('body','')))}",styles["BodyText"]))
    story += [Spacer(1,8),Paragraph("Recommendations",styles["Heading2"])]
    for r in recs: story.append(Paragraph(f"<b>{html.escape(str(r.get('priority','Info')))}</b> — <b>{html.escape(str(r.get('title','')))}</b>: {html.escape(str(r.get('body','')))}",styles["BodyText"]))
    story += [PageBreak(),Paragraph("Cleaned dataset preview",styles["Heading2"])]
    preview=clean_df.head(25).astype(str)
    if preview.shape[1] > 8:
        preview = preview.iloc[:, :8]
    data=[list(preview.columns)]+preview.values.tolist()
    if data:
        widths=[max(60,min(92,420/max(1,len(data[0])))) for _ in data[0]]
        pt=Table(data,colWidths=widths,repeatRows=1)
        pt.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#272A56")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTSIZE",(0,0),(-1,-1),6),("GRID",(0,0),(-1,-1),.15,colors.HexColor("#CBD5E1"))]))
        story.append(pt)
    doc.build(story); return out.getvalue()


def csv_bytes(df: pd.DataFrame) -> bytes: return df.to_csv(index=False).encode("utf-8-sig")
